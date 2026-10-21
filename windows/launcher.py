"""Windowed onedir entry point. No development Python, external cache or uploads on disk."""
import ctypes
from ctypes import wintypes
import io
import json
import multiprocessing
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
import traceback

TITLE="Anonimizador CEN-CINAI"

def root_path():
    return Path(getattr(sys,"_MEIPASS",Path(__file__).resolve().parents[1])).resolve()

def environment(root):
    # Covers startup as well as file processing. The existing model loader uses an
    # explicit bundle-relative directory and local_files_only, never this cache.
    os.environ.update(HF_HUB_OFFLINE="1",TRANSFORMERS_OFFLINE="1",HF_HUB_DISABLE_TELEMETRY="1",
        DO_NOT_TRACK="1",TOKENIZERS_PARALLELISM="false",STREAMLIT_BROWSER_GATHER_USAGE_STATS="false",
        HF_HOME=str(root/"models"/"_offline_cache"),HF_HUB_CACHE=str(root/"models"/"_offline_cache"/"hub"))
    if sys.stdout is None:sys.stdout=open(os.devnull,"w",encoding="utf-8")
    if sys.stderr is None:sys.stderr=open(os.devnull,"w",encoding="utf-8")
    from anonymizer.privacy import block_external_network
    block_external_network()

def check_resources(root):
    required=["app.py",".streamlit/config.toml","profiles/synthetic_demo_profile.yaml",
        "models/spanish-pii/model.safetensors","models/spanish-pii/config.json",
        "models/spanish-pii/tokenizer.json","models/spanish-pii/tokenizer_config.json"]
    if any(not (root/p).is_file() for p in required):raise RuntimeError("resources")

def free_port(start=8501):
    for port in range(start,start+100):
        try:
            with socket.socket() as s:
                if hasattr(socket,"SO_EXCLUSIVEADDRUSE"):s.setsockopt(socket.SOL_SOCKET,socket.SO_EXCLUSIVEADDRUSE,1)
                s.bind(("127.0.0.1",port))
            return port
        except OSError:continue
    raise RuntimeError("port")

def server(root,port):
    os.chdir(root)
    from streamlit.web import bootstrap
    flags={"server.address":"127.0.0.1","server.port":port,"server.headless":True,
        "browser.gatherUsageStats":False,"browser.serverAddress":"127.0.0.1",
        "server.fileWatcherType":"none","global.developmentMode":False,
        "client.showErrorDetails":"none"}
    bootstrap.load_config_options(flags)
    bootstrap.run(str(root/"app.py"),False,[],flags)

class WindowsJob:
    """Windows closes the entire server job even if the launcher is forcibly killed."""
    def __init__(self):
        class Basic(ctypes.Structure):
            _fields_=[("ProcessTime",ctypes.c_longlong),("JobTime",ctypes.c_longlong),("Flags",wintypes.DWORD),
                ("MinWorkingSet",ctypes.c_size_t),("MaxWorkingSet",ctypes.c_size_t),("ActiveProcesses",wintypes.DWORD),
                ("Affinity",ctypes.c_size_t),("Priority",wintypes.DWORD),("Scheduling",wintypes.DWORD)]
        class IO(ctypes.Structure):
            _fields_=[(n,ctypes.c_ulonglong) for n in ("ReadOps","WriteOps","OtherOps","ReadBytes","WriteBytes","OtherBytes")]
        class Extended(ctypes.Structure):
            _fields_=[("Basic",Basic),("IO",IO),("ProcessMemory",ctypes.c_size_t),("JobMemory",ctypes.c_size_t),
                ("PeakProcess",ctypes.c_size_t),("PeakJob",ctypes.c_size_t)]
        self.kernel=ctypes.WinDLL("kernel32",use_last_error=True)
        self.kernel.CreateJobObjectW.restype=wintypes.HANDLE
        self.kernel.CreateJobObjectW.argtypes=[ctypes.c_void_p,wintypes.LPCWSTR]
        self.kernel.SetInformationJobObject.argtypes=[wintypes.HANDLE,ctypes.c_int,ctypes.c_void_p,wintypes.DWORD]
        self.kernel.AssignProcessToJobObject.argtypes=[wintypes.HANDLE,wintypes.HANDLE]
        self.kernel.CloseHandle.argtypes=[wintypes.HANDLE]
        self.handle=self.kernel.CreateJobObjectW(None,None)
        limits=Extended();limits.Basic.Flags=0x2000 # JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        if not self.handle or not self.kernel.SetInformationJobObject(self.handle,9,ctypes.byref(limits),ctypes.sizeof(limits)):
            self.close();raise OSError("job")
    def assign(self,process):
        if not self.kernel.AssignProcessToJobObject(self.handle,wintypes.HANDLE(int(process._handle))):raise OSError("job assignment")
    def close(self):
        if getattr(self,"handle",None):self.kernel.CloseHandle(self.handle);self.handle=None

def mutex():
    kernel=ctypes.WinDLL("kernel32",use_last_error=True)
    kernel.CreateMutexW.restype=wintypes.HANDLE
    kernel.CreateMutexW.argtypes=[ctypes.c_void_p,wintypes.BOOL,wintypes.LPCWSTR]
    handle=kernel.CreateMutexW(None,False,"Local\\CENCINAI_Anonimizador_v1")
    if not handle:raise OSError("mutex")
    if ctypes.get_last_error()==183:
        user=ctypes.WinDLL("user32",use_last_error=True)
        user.FindWindowW.restype=wintypes.HWND;user.FindWindowW.argtypes=[wintypes.LPCWSTR,wintypes.LPCWSTR]
        user.ShowWindow.argtypes=[wintypes.HWND,ctypes.c_int]
        user.SetForegroundWindow.argtypes=[wintypes.HWND]
        hwnd=user.FindWindowW(None,TITLE)
        if hwnd:user.ShowWindow(hwnd,9);user.SetForegroundWindow(hwnd)
        else:message("La herramienta ya está iniciándose. Espere unos momentos.")
        kernel.CloseHandle(wintypes.HANDLE(handle))
        return None
    return handle

def message(text):
    if any(flag in sys.argv for flag in ("--qa-status","--qa-probe","--diagnostics-log")):return
    ctypes.windll.user32.MessageBoxW(None,text,TITLE,0x10)

def write_status(path,**info):
    if path:Path(path).write_text(json.dumps(info,indent=2),encoding="utf-8")

def open_browser(url,root):
    import webbrowser
    # External browsers must not inherit the bundle's DLL search directory.
    frozen=bool(getattr(sys,"frozen",False))
    if frozen:ctypes.windll.kernel32.SetDllDirectoryW(None)
    try:return webbrowser.open(url)
    finally:
        if frozen:ctypes.windll.kernel32.SetDllDirectoryW(str(root))

def qa_probe(root,target):
    """Build validation only: synthetic in-memory workflow, never a user entry."""
    from io import BytesIO
    from openpyxl import Workbook,load_workbook
    from anonymizer.spanish_ner import SpanishNER
    from anonymizer.service import Service
    from anonymizer.review import build_review,edit_row,confirm_review,COMPLETE
    ner=SpanishNER()
    assert str(ner.pipe.tokenizer.name_or_path).startswith(str(root))
    ner.batch(["Contacto ficticio: 8888-1234."])
    w=Workbook();w.active.title="Datos";w.active.append(["Nombre","Correo","Observaciones"])
    w.active.append(["Persona ficticia","qa@example.test","Nota ficticia"])
    b=BytesIO();w.save(b);data=b.getvalue();service=Service(ner)
    analysis=service.analyze(data,"qa.xlsx");rows=build_review(analysis.columns)
    rows=[edit_row(r,kind=k,selected=True,resolution="Anonimizar",**({'method':COMPLETE} if k=="Texto libre" else {})) for r,k in zip(rows,["Persona","Correo","Texto libre"])]
    plan=confirm_review(rows,analysis.columns,data,{"Datos":1},True)
    result=service.anonymize_reviewed(data,"qa.xlsx",plan)
    out=load_workbook(BytesIO(result.data)).active
    assert out['A2'].value!="Persona ficticia" and out['B2'].value!="qa@example.test" and out['C2'].value=="[CONTENIDO ANONIMIZADO]"
    denied=False
    try:socket.create_connection(("192.0.2.1",80),timeout=.1)
    except PermissionError:denied=True
    assert denied
    write_status(target,result="PASS",frozen=bool(getattr(sys,"frozen",False)),executable=sys.executable,
        bundle_root=str(root),tokenizer_path=str(ner.pipe.tokenizer.name_or_path),external_connection_denied=denied,
        actual_model_inference=True,synthetic_workflow=True,cells_changed=result.cells,
        console_window=bool(ctypes.windll.kernel32.GetConsoleWindow()),sys_path=sys.path)

def desktop(root,status_path=None,diagnostic_log=None):
    import tkinter as tk
    from tkinter import ttk
    import urllib.request
    held_mutex=mutex()
    if held_mutex is None:return 0
    window=tk.Tk();window.title(TITLE);window.geometry("550x300");window.resizable(False,False)
    window.configure(bg="#F5F1F5")
    window.iconbitmap(str(root/"assets"/"anonimizador.ico"))
    tk.Label(window,text=TITLE,font=("Segoe UI",19,"bold"),fg="#663366",bg="#F5F1F5").pack(pady=(25,12))
    label=tk.Label(window,text="La herramienta está iniciando…",font=("Segoe UI",12),bg="#F5F1F5",wraplength=500)
    label.pack(pady=8)
    progress=ttk.Progressbar(window,mode="indeterminate",length=450);progress.pack(pady=12);progress.start()
    tk.Label(window,text="Procesamiento local · Espere la apertura del navegador.\nMantenga esta ventana abierta mientras utiliza la herramienta.",font=("Segoe UI",10),fg="#454045",bg="#F5F1F5").pack(pady=6)
    controls=tk.Frame(window,bg="#F5F1F5");controls.pack(pady=12)
    port=free_port();url=f"http://127.0.0.1:{port}"
    reopen=ttk.Button(controls,text="Abrir aplicación",command=lambda:open_browser(url,root),state="disabled");reopen.pack(side="left",padx=8)
    process=None;job=None;started=time.monotonic();ready=False
    local_http=urllib.request.build_opener(urllib.request.ProxyHandler({}))
    def close():
        if job:job.close()
        if process:
            try:process.wait(timeout=5)
            except subprocess.TimeoutExpired:process.kill();process.wait(timeout=5)
        window.destroy()
    ttk.Button(controls,text="Cerrar herramienta",command=close).pack(side="left",padx=8)
    window.protocol("WM_DELETE_WINDOW",close)
    def start():
        nonlocal process,job
        try:
            job=WindowsJob()
            cmd=[sys.executable,"--server",str(port)] if getattr(sys,"frozen",False) else [sys.executable,str(Path(__file__).resolve()),"--server",str(port)]
            if diagnostic_log:cmd.extend(["--diagnostics-log",diagnostic_log])
            process=subprocess.Popen(cmd,cwd=root,creationflags=subprocess.CREATE_NO_WINDOW)
            job.assign(process)
            window.after(400,poll)
        except Exception:
            if diagnostic_log:Path(diagnostic_log).write_text(traceback.format_exc(),encoding="utf-8")
            if process and process.poll() is None:process.terminate()
            message("No fue posible iniciar la herramienta. Cierre otra instancia y compruebe que copió la carpeta completa.");close()
    def poll():
        nonlocal ready
        if process.poll() is not None:
            message("La aplicación se detuvo. Compruebe que la carpeta está completa y vuelva a abrirla.");close();return
        if not ready:
            try:
                with local_http.open(url+"/_stcore/health",timeout=.2) as response:
                    if response.status==200:
                        ready=True;progress.stop();progress.pack_forget();reopen.configure(state="normal")
                        opened=open_browser(url,root)
                        label.configure(text="Aplicación disponible en su navegador.\nEl primer análisis puede tardar mientras carga el modelo.")
                        write_status(status_path,ready=True,url=url,pid=os.getpid(),server_pid=process.pid,
                            browser_auto_open=bool(opened),frozen=bool(getattr(sys,"frozen",False)),
                            executable=sys.executable,bundle_root=str(root),console_window=bool(ctypes.windll.kernel32.GetConsoleWindow()),
                            external_connections_blocked=True,hf_home=os.environ["HF_HOME"],sys_path=sys.path)
            except OSError:pass
            if not ready and time.monotonic()-started>180:
                message("El inicio está tardando demasiado. Cierre aplicaciones para liberar memoria y vuelva a intentarlo.");close();return
        window.after(700,poll)
    try:window.after(150,start);window.mainloop()
    finally:
        if job:job.close()
        ctypes.windll.kernel32.CloseHandle(wintypes.HANDLE(held_mutex))
    return 0

def main():
    multiprocessing.freeze_support()
    root=root_path();environment(root)
    args=sys.argv[1:]
    def argument(flag):return args[args.index(flag)+1] if flag in args else None
    try:
        check_resources(root)
        if "--qa-probe" in args:qa_probe(root,argument("--qa-probe"));return 0
        if "--server" in args:server(root,int(argument("--server")));return 0
        return desktop(root,argument("--qa-status"),argument("--diagnostics-log"))
    except BaseException:
        if "--diagnostics-log" in args:Path(argument("--diagnostics-log")).write_text(traceback.format_exc(),encoding="utf-8")
        if "--server" not in args:message("No fue posible abrir la herramienta. Copie la carpeta completa y vuelva a intentarlo.")
        return 1

if __name__=="__main__":sys.exit(main())
