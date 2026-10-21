from io import BytesIO
from pathlib import PurePath
import re
import warnings
from zipfile import ZipFile, BadZipFile

from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.packaging.core import DocumentProperties

MAX_BYTES=50*1024**2
MAX_UNCOMPRESSED=250*1024**2
MAX_CELLS=2_000_000

class WorkbookError(ValueError): pass

def load(data,filename):
    if PurePath(filename).suffix.lower()!=".xlsx":
        raise WorkbookError("Solo se admiten archivos .xlsx. Guarde una copia en ese formato desde Excel.")
    if not data or len(data)>MAX_BYTES:
        raise WorkbookError("El archivo está vacío o supera el límite de 50 MiB.")
    try:
        with ZipFile(BytesIO(data)) as z:
            entries=z.infolist()
            if len(entries)>10000 or sum(x.file_size for x in entries)>MAX_UNCOMPRESSED:
                raise WorkbookError("El contenido descomprimido excede los límites admitidos.")
            names=[x.filename for x in entries]
            if len(names)!=len(set(names)): raise WorkbookError("El archivo contiene entradas duplicadas no admitidas.")
            if "xl/workbook.xml" not in names: raise WorkbookError("El archivo seleccionado no parece ser un Excel válido.")
            blocked=("vbaproject","externallinks/","connections","querytables/","pivot","embeddings/","drawings/","media/","comments","threadedcomments","persons/","customxml/","customproperties","custom.xml","activex/","ctrlprops/","printersettings/","tables/")
            if any(any(b in name.casefold() for b in blocked) for name in names):
                raise WorkbookError("El archivo contiene objetos, comentarios, macros o contenido externo que esta versión no puede procesar.")
            for name in names:
                if name.endswith(".rels"):
                    root=fromstring(z.read(name))
                    if any(e.get("TargetMode","").lower()=="external" for e in root):
                        raise WorkbookError("El archivo contiene vínculos externos no admitidos.")
            root=fromstring(z.read("xl/workbook.xml"))
            if any(e.tag.endswith("workbookProtection") and any(v not in ("0","false","") for v in e.attrib.values()) for e in root.iter()):
                raise WorkbookError("El libro está protegido. Solicite una copia autorizada sin protección.")
            # Bound dimensions before openpyxl allocates cell objects.
            count=0
            for name in names:
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml",name):
                    sheet=fromstring(z.read(name)); count+=sum(1 for e in sheet.iter() if e.tag.endswith("}c"))
                    if count>MAX_CELLS: raise WorkbookError("El libro supera el límite de 2 millones de celdas.")
                    for e in sheet.iter():
                        if e.tag.endswith("}extLst") or e.tag.endswith("}legacyDrawing"):
                            raise WorkbookError("El archivo contiene extensiones de Excel no admitidas.")
                        if e.tag.endswith("}filterColumn"):
                            raise WorkbookError("El libro contiene criterios de filtro guardados. Quite los criterios en una copia antes de procesarla.")
                        if e.tag.endswith("}formula") and e.text and ('"' in e.text or '[' in e.text or re.search(r'\d{8,}|\bDATE\s*\(\s*\d',e.text,re.I)):
                            raise WorkbookError("El libro contiene reglas de formato condicional con texto no admitido.")
            if "xl/styles.xml" in names:
                styles=fromstring(z.read("xl/styles.xml"))
                for e in styles.iter():
                    if e.tag.endswith("}numFmt"):
                        literals=re.findall(r'"([^"]*)"',e.get("formatCode",""))
                        if any(x not in ("$","₡","€","USD","CRC","%"," "," de ","/","-",":") for x in literals):
                            raise WorkbookError("El libro contiene formatos numéricos con texto no admitido.")
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            wb=load_workbook(BytesIO(data),data_only=False,keep_links=False)
        if len(wb.worksheets)>50: raise WorkbookError("El libro supera el límite de 50 hojas.")
        if wb.defined_names: raise WorkbookError("El libro contiene nombres definidos personalizados no admitidos.")
        cells=0
        for ws in wb:
            if ws.defined_names: raise WorkbookError("Una hoja contiene nombres definidos personalizados no admitidos.")
            if ws.max_row>100000 or ws.max_column>500: raise WorkbookError("Una hoja supera 100.000 filas o 500 columnas.")
            cells+=ws.max_row*ws.max_column
            if cells>MAX_CELLS: raise WorkbookError("El libro supera el límite de 2 millones de celdas rectangulares.")
            if ws.protection.sheet: raise WorkbookError("Una hoja está protegida. Solicite una copia autorizada sin protección.")
            if any(str(x) for x in [ws.oddHeader,ws.evenHeader,ws.firstHeader,ws.oddFooter,ws.evenFooter,ws.firstFooter]):
                raise WorkbookError("El libro contiene encabezados o pies de impresión con texto no admitido.")
            for row in ws:
                for c in row:
                    if c.comment or c.hyperlink: raise WorkbookError("El libro contiene comentarios o vínculos no admitidos.")
                    if c.data_type=="f" and (not isinstance(c.value,str) or any(x in c.value for x in ['"','[',']']) or re.search(r"\b(HYPERLINK|WEBSERVICE|IMAGE|RTD|DDE)\s*\(",c.value,re.I)):
                        raise WorkbookError("El libro contiene fórmulas con texto, contenido externo o estructuras no admitidas.")
                    if c.data_type=="f" and (not re.search(r"\$?[A-Z]{1,3}\$?\d+",c.value,re.I) or re.search(r"\d{7,}|\bDATE\s*\(\s*\d",c.value,re.I)):
                        raise WorkbookError("El libro contiene fórmulas constantes que podrían conservar identificaciones o fechas personales. Prepare una copia sin esas fórmulas.")
            for dv in ws.data_validations.dataValidation:
                if dv.prompt or dv.error or dv.promptTitle or dv.errorTitle or any('"' in str(x) or '[' in str(x) or re.search(r'\d{8,}|\bDATE\s*\(\s*\d',str(x),re.I) for x in [dv.formula1,dv.formula2]):
                    raise WorkbookError("El libro contiene validaciones con texto o referencias externas no admitidas.")
        return wb
    except WorkbookError: raise
    except Exception:
        raise WorkbookError("No fue posible abrir el archivo. El archivo seleccionado no parece ser un Excel válido o contiene características no admitidas.") from None

def headers_for(wb,header_rows=None):
    result={}
    for ws in wb:
        row=int((header_rows or {}).get(ws.title,1))
        if row<1 or row>ws.max_row: raise WorkbookError("La fila de encabezados no existe en una hoja.")
        if any(c.value is not None for values in ws.iter_rows(max_row=row-1) for c in values) if row>1 else False:
            raise WorkbookError("Hay contenido por encima del encabezado. Prepare una copia con una tabla por hoja.")
        result[ws.title]=row
    return result

def export(wb):
    wb.properties=DocumentProperties(creator="CEN-CINAI Anonimizador",lastModifiedBy=None)
    if hasattr(wb,"calculation") and wb.calculation:
        wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True
    out=BytesIO(); wb.save(out); return out.getvalue()

def output_name(filename):
    # Strip both path separators even on non-Windows development systems.
    stem=re.split(r"[/\\]",filename)[-1][:-5]
    stem=re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",stem)[:150]
    return (stem or "archivo")+"_anonimizado.xlsx"
