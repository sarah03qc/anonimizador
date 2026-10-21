"""Synthetic H01–H25: human authority and complete structured replacement."""
from dataclasses import replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
import hashlib
import pytest
from openpyxl import Workbook,load_workbook
from presidio_analyzer import RecognizerResult
from anonymizer.models import Column
from anonymizer.review import *
from anonymizer.profiles import load_profile,apply_profile,suggest_profiles
from anonymizer.service import Service

ROOT=Path(__file__).resolve().parents[1]
class NoNER:
    def batch(self,texts):raise AssertionError("Structured/full-cell replacement must not invoke NER")
class MissingNER:
    def batch(self,texts):return [[] for _ in texts]

def book(headers,values,sheets=("Datos",)):
    w=Workbook();w.remove(w.active)
    for name in sheets:
        s=w.create_sheet(name);s.append(headers)
        for row in values:s.append(row)
    b=BytesIO();w.save(b);return b.getvalue()

def col(index=1,header="campo_x",kind="No determinado",state="No detectada como sensible",selected=False,sheet="Datos"):
    return Column(sheet,index,header,state,kind,"Baja",selected,rows=3)

def run(data,columns,rows,ner=None):
    headers={c.sheet:1 for c in columns}
    plan=confirm_review(rows,columns,data,headers,attested=True)
    return Service(ner or NoNER()).anonymize_reviewed(data,"ficticio.xlsx",plan)

def test_H01_all_columns_all_sheets():
    columns=[col(1),col(2),col(1,sheet="Oculta")]
    assert [r.key for r in build_review(columns)]==[(c.sheet,c.index) for c in columns]

def test_H02_sensitive_preselected():
    r=build_review([col(kind="Persona",state="Sensible",selected=True)])[0]
    assert r.selected and automatic_label(r)=="🔴 Sensible"

def test_H03_non_detected_visible_selectable():
    rows=build_review([col()]);assert visible_rows(rows,"Todas")==rows
    assert "Sin indicadores detectados" in editor_records(rows)[0]["Detección"]
    assert not rows[0].selected

def test_H04_manual_green_selection():
    r=edit_row(build_review([col()])[0],kind="Persona",selected=True)
    assert r.selected and not validation_errors([r])

def test_H05_manual_red_deselection():
    r=edit_row(build_review([col(kind="Persona",state="Sensible",selected=True)])[0],selected=False)
    assert not r.selected and r.resolution==PRESERVE

def test_H06_type_override():
    r=edit_row(build_review([col(kind="Persona",state="Sensible",selected=True)])[0],kind="Correo")
    assert r.kind=="Correo" and r.method==STRUCTURED

def test_H07_yellow_requires_explicit_decision_even_hidden():
    rows=build_review([col(1,kind="Persona",state="Revisar",selected=True),col(2,kind="Persona",state="Sensible",selected=True)])
    assert validation_errors(rows)
    assert len(visible_rows(rows,"Sin indicadores"))==0
    with pytest.raises(ReviewError):confirm_review(rows,[col(1),col(2)],b"x",{"Datos":1},True)
    for decision in ("Anonimizar",PRESERVE):
        fixed=[edit_row(rows[0],resolution=decision),rows[1]]
        assert not validation_errors(fixed)

def test_H08_no_confirmation_no_run():
    data=book(["Nombre"],[["Dato ficticio"]]);columns=[col(kind="Persona",state="Sensible",selected=True)]
    with pytest.raises(ReviewError):Service(NoNER()).anonymize_reviewed(data,"x.xlsx",None)
    with pytest.raises(ReviewError):confirm_review(build_review(columns),columns,data,{"Datos":1})

@pytest.mark.parametrize("case,kind,value",[("H09","Persona","Qzxv sin patrón"),("H10","Identificación","ID peculiar"),("H11","Teléfono",123),("H12","Correo","sin-arroba"),("H13","Fecha personal",datetime(1988,4,3)),("H14","Dirección / ubicación","x"),("financial","Financiero",450),("credential","Credencial / usuario","u"),("other","Otro sensible","z")])
def test_H09_H14_all_nonempty_without_NER(case,kind,value):
    data=book(["campo_a"],[[value],[value],[None],["  "]])
    columns=[col()];rows=[edit_row(build_review(columns)[0],kind=kind,selected=True)]
    result=run(data,columns,rows);out=load_workbook(BytesIO(result.data)).active
    assert out["A2"].value!=value and out["A3"].value!=value
    assert out["A2"].value==out["A3"].value and result.cells==2
    assert out["A4"].value is None and out["A5"].value=="  "
    if kind=="Fecha personal":assert isinstance(out["A2"].value,datetime)

def test_H15_preserve_values_and_style():
    data=book(["campo_a","Región"],[["Nombre ficticio","Central ficticia"],["Otro",12]])
    columns=[col(1),col(2,header="Región",kind="Institucional")]
    rows=build_review(columns);rows[0]=edit_row(rows[0],kind="Persona",selected=True)
    out=load_workbook(BytesIO(run(data,columns,rows).data)).active
    before=load_workbook(BytesIO(data)).active
    assert [(c.value,c.data_type,c.style_id) for c in out["B"]]==[(c.value,c.data_type,c.style_id) for c in before["B"]]

def test_H16_empty_complete_mode():
    data=book(["Observaciones"],[[None],["  "],[""],["Dato ficticio"]])
    columns=[col()];rows=[edit_row(build_review(columns)[0],kind="Texto libre",method=COMPLETE,selected=True)]
    out=load_workbook(BytesIO(run(data,columns,rows).data)).active
    assert [out.cell(i,1).value for i in range(2,6)]==[None,"  ",None,"[CONTENIDO ANONIMIZADO]"]

def test_H17_entity_mode_calls_detector_preserves_context():
    class EntityNER:
        calls=0
        def batch(self,texts):
            self.calls+=1
            return [[RecognizerResult("PERSON",6,16,.99)] for t in texts]
    ner=EntityNER();data=book(["Observaciones"],[["Llamó Ana Prueba ayer."]]);columns=[col()]
    rows=[edit_row(build_review(columns)[0],kind="Texto libre",method=ENTITIES,selected=True)]
    value=load_workbook(BytesIO(run(data,columns,rows,ner).data)).active["A2"].value
    assert ner.calls and value.startswith("Llamó ") and value.endswith(" ayer.") and "Ana Prueba" not in value

def test_H18_full_cell_without_NER():
    data=book(["Observaciones"],[["Contenido sintético que no reconoce ningún detector"]]);columns=[col()]
    rows=[edit_row(build_review(columns)[0],kind="Texto libre",method=COMPLETE,selected=True)]
    result=run(data,columns,rows)
    assert load_workbook(BytesIO(result.data)).active["A2"].value=="[CONTENIDO ANONIMIZADO]"
    assert result.counts=={"Texto libre — contenido completo":1}

def test_empty_entity_mode_never_fills_or_detects_blank_cells():
    data=book(["Observaciones"],[[None],["  "],[""]]);columns=[col()]
    rows=[edit_row(build_review(columns)[0],kind="Texto libre",method=ENTITIES,selected=True)]
    result=run(data,columns,rows)
    out=load_workbook(BytesIO(result.data)).active
    assert [out.cell(i,1).value for i in range(2,5)]==[None,"  ",None]
    assert result.cells==0

def test_H19_missed_column_human_override():
    data=book(["campo_x"],[["N."]]);analysis=Service(MissingNER()).analyze(data,"x.xlsx")
    assert not analysis.columns[0].selected
    rows=[edit_row(build_review(analysis.columns)[0],kind="Persona",selected=True)]
    assert load_workbook(BytesIO(run(data,analysis.columns,rows).data)).active["A2"].value!="N."

def test_H20_false_positive_human_preserves():
    data=book(["Establecimiento","Nombre"],[["CEN ficticio","Persona ficticia"]])
    columns=[col(1,"Establecimiento","Persona","Sensible",True),col(2,"Nombre","Persona","Sensible",True)]
    rows=build_review(columns);rows[0]=edit_row(rows[0],kind="Conservar",selected=False)
    out=load_workbook(BytesIO(run(data,columns,rows).data)).active
    assert out["A2"].value=="CEN ficticio" and out["B2"].value!="Persona ficticia"

def test_H21_consistent_across_sheets():
    data=book(["Nombre"],[["Persona ficticia"]],("Datos","Otra"))
    columns=[col(kind="Persona",state="Sensible",selected=True,sheet=s) for s in ("Datos","Otra")]
    out=load_workbook(BytesIO(run(data,columns,build_review(columns)).data))
    assert out["Datos"]["A2"].value==out["Otra"]["A2"].value

def test_H22_profile_defaults_and_suggestion():
    profile=load_profile(ROOT/"profiles/synthetic_demo_profile.yaml")
    columns=[col(1,"campo_a"),col(2,"campo_b"),col(3,"Región","Institucional")]
    rows=apply_profile(build_review(columns),profile)
    assert profile.name=="SYNTHETIC_DEMO_PROFILE" and suggest_profiles(columns,[profile])==[profile]
    assert rows[0].selected and rows[0].kind=="Persona" and rows[1].kind=="Identificación" and not rows[2].selected

def test_H23_human_overrides_profile():
    profile=load_profile(ROOT/"profiles/synthetic_demo_profile.yaml")
    r=apply_profile(build_review([col(1,"campo_a")]),profile)[0]
    assert not edit_row(r,kind="Conservar").selected
    assert edit_row(r,kind="Correo").kind=="Correo"

def test_H24_unknown_headers_complete_after_review():
    data=(ROOT/"test_data/human_review/unknown_headers.xlsx").read_bytes()
    analysis=Service(MissingNER()).analyze(data,"x.xlsx")
    rows=build_review(analysis.columns)
    for i,kind in enumerate(("Persona","Identificación","Teléfono","Correo")):
        rows[i]=edit_row(rows[i],kind=kind,selected=True,resolution="Anonimizar")
    # Resolve every other yellow explicitly, preserving the non-personal fields.
    rows=[edit_row(r,resolution=PRESERVE) if r.index>4 else r for r in rows]
    out=load_workbook(BytesIO(run(data,analysis.columns,rows).data)).active
    original=load_workbook(BytesIO(data)).active
    for row in original.iter_rows(min_row=2):
        for c in row:
            value=out.cell(c.row,c.column).value
            assert (value!=c.value) if c.column<=4 and c.value is not None and str(c.value).strip() else value==c.value

def test_H25_original_intact():
    data=book(["Nombre"],[["Persona ficticia"]]);before=hashlib.sha256(data).hexdigest()
    columns=[col(kind="Persona",state="Sensible",selected=True)]
    run(data,columns,build_review(columns))
    assert hashlib.sha256(data).hexdigest()==before
    assert load_workbook(BytesIO(data)).active["A2"].value=="Persona ficticia"

def test_confirmation_bound_to_file_and_all_columns():
    data=book(["Nombre","Otro"],[["Persona ficticia",123]])
    columns=[col(1,kind="Persona",state="Sensible",selected=True),col(2)]
    rows=build_review(columns);plan=confirm_review(rows,columns,data,{"Datos":1},True)
    with pytest.raises(ReviewError):Service(NoNER()).anonymize_reviewed(data+b"x","x.xlsx",plan)
    with pytest.raises(ReviewError):Service(NoNER()).anonymize_reviewed(data,"x.xlsx",replace(plan,decisions=plan.decisions[:1]))
    with pytest.raises(ReviewError):confirm_review(rows[:1],columns,data,{"Datos":1},True)

@pytest.mark.parametrize("kind,method",[("No determinado",STRUCTURED),("Texto libre",STRUCTURED),("Persona",ENTITIES),("Persona",PRESERVE)])
def test_invalid_type_method_blocked(kind,method):
    row=edit_row(build_review([col()])[0],kind=kind,method=method,selected=True)
    assert validation_errors([row])

def test_yellow_profile_does_not_confirm():
    p=load_profile(ROOT/"profiles/synthetic_demo_profile.yaml")
    row=apply_profile(build_review([col(1,"campo_a",state="Revisar")]),p)[0]
    assert row.selected and row.resolution==PENDING and validation_errors([row])

def test_profile_safe_load_rejects_code(tmp_path):
    p=tmp_path/"bad.yaml";p.write_text("!!python/object/apply:os.system ['echo unsafe']")
    with pytest.raises(ReviewError):load_profile(p)

def test_editor_merges_human_choices():
    row=build_review([col()])[0];record=editor_records([row])[0]
    record.update({"Anonimizar":True,"Tipo":"Persona"})
    changed=merge_editor(row,record)
    assert changed.selected and changed.resolution=="Anonimizar" and changed.method==STRUCTURED
