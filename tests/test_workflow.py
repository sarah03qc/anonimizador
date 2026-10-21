"""Privacy and preservation regression tests; fixtures are exclusively fictitious."""
import hashlib
from dataclasses import replace
from datetime import datetime
from copy import copy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]

@pytest.fixture(scope="session")
def service():
    from anonymizer.service import Service
    from anonymizer.spanish_ner import SpanishNER
    return Service(SpanishNER())

def book(headers=None, rows=None):
    w=Workbook(); s=w.active; s.title="Datos"
    s.append(headers or ["Nombre completo","Correo","Monto"])
    for row in rows or [["María José Peña Quesada","demo@example.test",125000]]: s.append(row)
    b=BytesIO(); w.save(b); return b.getvalue()

def test_T01_valid(service):
    assert service.analyze(book(),"prueba.xlsx").columns

def test_T02_corrupt(service):
    from anonymizer.workbook import WorkbookError
    with pytest.raises(WorkbookError): service.analyze(b"corrupto","prueba.xlsx")

def test_T03_T21_T31_original_unchanged(service):
    data=book(); before=hashlib.sha256(data).digest()
    a=service.analyze(data,"prueba.xlsx")
    service.anonymize(data,"prueba.xlsx",{("Datos",1):"Persona"})
    assert hashlib.sha256(data).digest()==before
    assert load_workbook(BytesIO(data)).active["A2"].value=="María José Peña Quesada"

def test_T04_T05_T06_T23_T24_T25_T26_T27_T28_T29(service):
    data=(ROOT/"test_data/complejo.xlsx").read_bytes()
    result=service.anonymize(data,"complejo.xlsx",{("Datos",1):"Persona",("Otra hoja",1):"Persona",("Oculta",1):"Correo"})
    w=load_workbook(BytesIO(result.data)); original=load_workbook(BytesIO(data))
    assert w.sheetnames==["Datos","Otra hoja","Oculta"]
    for s in w:
        o=original[s.title]
        assert [c.value for c in s[1]]==[c.value for c in o[1]]
        assert s.sheet_state==o.sheet_state
        assert list(s.merged_cells.ranges)==list(o.merged_cells.ranges)
        assert s.freeze_panes==o.freeze_panes
        assert s.column_dimensions["A"].width==o.column_dimensions["A"].width
        assert copy(s["A1"].font)==copy(o["A1"].font)
        assert copy(s["A1"].fill)==copy(o["A1"].fill)
        assert s.row_dimensions[1].height==o.row_dimensions[1].height
        assert s.auto_filter.ref==o.auto_filter.ref
    assert w["Datos"]["D2"].value=="=B2*2"
    assert w["Datos"]["B2"].value==125000
    assert isinstance(w["Datos"]["B2"].value,int)
    assert isinstance(w["Datos"]["C2"].value,datetime)
    assert len(w["Datos"].data_validations.dataValidation)==1
    assert w["Datos"]["A2"].value==w["Otra hoja"]["A2"].value

@pytest.mark.parametrize("header,kind",[("nombreCompleto","Persona"),("correo electrónico","Correo"),("contacto_1","Teléfono"),("fechaDeNacimiento","Fecha personal"),("DIMEX","Identificación"),("cuenta bancaria","Financiero"),("usuario","Credencial")])
def test_T07_T08_T09_header(service,header,kind):
    a=service.analyze(book([header],[["demo@example.test"]]),"prueba.xlsx")
    assert a.columns[0].kind==kind
    assert a.columns[0].state=="Sensible"

def test_T10_text(service):
    text="La madre María Rodríguez llamó al 8888-1234 y escribió a demo@example.test."
    data=book(["Observaciones"],[[text]])
    a=service.analyze(data,"prueba.xlsx"); assert a.columns[0].selected
    r=service.anonymize(data,"prueba.xlsx",{("Datos",1):"Texto con PII"})
    out=load_workbook(BytesIO(r.data)).active["A2"].value
    assert "8888-1234" not in out and "demo@example.test" not in out
    assert "La madre " in out and " llamó al " in out

@pytest.mark.parametrize("header,value",[("Código establecimiento",1100102),("Código OL",101),("Código Región",1),("Región","Central"),("Oficina Local","Oficina ficticia"),("Establecimiento","CEN de prueba"),("Año",2026),("Monto",125000),("Edad",7),("id",123456789),("numero",123456789)])
def test_T14_T15_T16_T17_false_positives(service,header,value):
    c=service.analyze(book([header],[[value]]),"prueba.xlsx").columns[0]
    assert c.state!="Sensible"
    assert not c.selected

def test_T18_T22_T32_manual_selection(service):
    data=book(["codigo","Monto"],[[1100102,125000]])
    r=service.anonymize(data,"prueba.xlsx",{("Datos",1):"Otro"})
    w=load_workbook(BytesIO(r.data))
    assert w.active["A2"].value!=1100102
    assert w.active["B2"].value==125000
    assert r.filename=="prueba_anonimizado.xlsx"

def test_T19_consistency(service):
    data=book(["Nombre"],[["María José"],["  MARÍA  JOSÉ "]])
    r=service.anonymize(data,"x.xlsx",{("Datos",1):"Persona"})
    s=load_workbook(BytesIO(r.data)).active
    assert s["A2"].value==s["A3"].value

def test_T20_T30_no_persistence_no_logs(service,tmp_path,monkeypatch,caplog):
    monkeypatch.chdir(tmp_path)
    r=service.anonymize(book(),"x.xlsx",{("Datos",1):"Persona",("Datos",2):"Correo"})
    assert not list(tmp_path.iterdir())
    assert "María" not in caplog.text and "demo@example.test" not in caplog.text

@pytest.mark.parametrize("ext",["xls","xlsm","ods","csv"])
def test_unsupported_format(service,ext):
    from anonymizer.workbook import WorkbookError
    with pytest.raises(WorkbookError): service.analyze(book(),"x."+ext)

def test_protected(service):
    from anonymizer.workbook import WorkbookError
    w=load_workbook(BytesIO(book())); w.active.protection.sheet=True; b=BytesIO(); w.save(b)
    with pytest.raises(WorkbookError): service.analyze(b.getvalue(),"x.xlsx")

@pytest.mark.parametrize("feature",["comment","hyperlink","formula_literal","metadata_header"])
def test_unsupported_pii_surfaces(service,feature):
    from anonymizer.workbook import WorkbookError
    w=load_workbook(BytesIO(book())); s=w.active
    if feature=="comment": s["B2"].comment=Comment("texto privado","autor")
    if feature=="hyperlink": s["B2"].hyperlink="https://example.test/privado"
    if feature=="formula_literal": s["C2"]='="texto privado"'
    if feature=="metadata_header": s.oddHeader.center.text="texto privado"
    b=BytesIO(); w.save(b)
    with pytest.raises(WorkbookError): service.analyze(b.getvalue(),"x.xlsx")

def test_metadata_cleaned(service):
    w=load_workbook(BytesIO(book())); w.properties.creator="autor ficticio"; w.properties.description="privado"
    b=BytesIO(); w.save(b)
    r=service.anonymize(b.getvalue(),"x.xlsx",{("Datos",1):"Persona"})
    props=load_workbook(BytesIO(r.data)).properties
    assert props.creator=="CEN-CINAI Anonimizador" and props.description is None

def test_selected_originals_absent_from_xlsx_package(service):
    data=book()
    r=service.anonymize(data,"x.xlsx",{("Datos",1):"Persona",("Datos",2):"Correo"})
    with ZipFile(BytesIO(r.data)) as package:
        xml=b" ".join(package.read(n) for n in package.namelist() if n.endswith(".xml"))
    assert b"demo@example.test" not in xml
    assert "María José Peña Quesada".encode() not in xml

def test_mapping_cleared_on_failure(service,monkeypatch):
    from anonymizer import service as module
    from anonymizer.replacement import Replacer
    instance=Replacer()
    monkeypatch.setattr(module,"Replacer",lambda:instance)
    def fail(_): raise RuntimeError("fallo simulado")
    monkeypatch.setattr(module.workbook,"export",fail)
    with pytest.raises(RuntimeError): service.anonymize(book(),"x.xlsx",{("Datos",1):"Persona"})
    assert instance.mapping.values=={} and instance.known=={}
