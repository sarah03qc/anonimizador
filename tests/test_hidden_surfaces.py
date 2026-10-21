from io import BytesIO
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.filters import FilterColumn,Filters
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from anonymizer.workbook import load,WorkbookError

def encoded(w):
    b=BytesIO(); w.save(b); return b.getvalue()

@pytest.mark.parametrize("feature",["filter","number_format","conditional_format"])
def test_rejects_literals_outside_selected_cells(feature):
    w=Workbook(); s=w.active; s.append(["Monto"]); s.append([123])
    if feature=="filter":
        s.auto_filter.ref="A1:A2"; s.auto_filter.filterColumn.append(FilterColumn(colId=0,filters=Filters(filter=["Privado Ficticio"])))
    elif feature=="number_format": s["A2"].number_format='0 "Privado Ficticio"'
    else: s.conditional_formatting.add("A2",FormulaRule(formula=['A2="Privado Ficticio"']))
    with pytest.raises(WorkbookError): load(encoded(w),"x.xlsx")

def test_empty_workbook_protection_does_not_reject():
    w=Workbook(); w.active.append(["Nombre"]); w.active.append(["Persona Ficticia"])
    assert load(encoded(w),"x.xlsx").active["A2"].value=="Persona Ficticia"

@pytest.mark.parametrize("formula",["=123456789","=DATE(1990,1,1)","=DATE(1990,1,1)+B2","=123456789+B2"])
def test_rejects_constant_formula_pii(formula):
    w=Workbook(); w.active.append(["Identificación"]); w.active.append([formula])
    with pytest.raises(WorkbookError):load(encoded(w),"x.xlsx")

@pytest.mark.parametrize("surface",["local_name","validation","conditional"])
def test_rejects_local_names_and_numeric_metadata(surface):
    w=Workbook(); s=w.active; s.append(["Monto"]); s.append([123])
    if surface=="local_name": s.defined_names.add(DefinedName("CampoPrivado",attr_text='"demo@example.test"',localSheetId=0))
    if surface=="validation":
        dv=DataValidation(type="whole",operator="equal",formula1="123456789"); dv.add("A2"); s.add_data_validation(dv)
    if surface=="conditional": s.conditional_formatting.add("A2",FormulaRule(formula=["A2=123456789"]))
    with pytest.raises(WorkbookError):load(encoded(w),"x.xlsx")
