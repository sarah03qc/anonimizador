import sys
from io import BytesIO
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from anonymizer.privacy import block_external_network
block_external_network()
from anonymizer.spanish_ner import SpanishNER
from anonymizer.service import Service
from openpyxl import Workbook,load_workbook

def main():
    w=Workbook(); w.active.append(["Nombre completo","Correo"]); w.active.append(["Persona Ficticia","demo@example.test"])
    b=BytesIO(); w.save(b); data=b.getvalue()
    service=Service(SpanishNER())
    analysis=service.analyze(data,"prueba.xlsx")
    assert analysis.columns[0].selected
    result=service.anonymize(data,"prueba.xlsx",{("Sheet",1):"Persona",("Sheet",2):"Correo"})
    assert load_workbook(BytesIO(result.data)).active["B2"].value!="demo@example.test"
    print("PASS: modelo, Presidio, Excel, sustitucion y bloqueo de conexiones externas.")

if __name__=="__main__": main()
