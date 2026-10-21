"""Generate only fictitious XLSX fixtures, never institutional records."""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "test_data"

def save(wb, name):
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32
        for c in ws[1]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = PatternFill("solid", fgColor="663366")
            c.alignment = Alignment(wrap_text=True)
            ws.column_dimensions[c.column_letter].width = 23
    wb.save(OUT / name)

def generate():
    OUT.mkdir(exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Ejemplo simple"
    ws.append(["Nombre completo", "Correo", "Teléfono", "Fecha nacimiento"])
    for i in range(100):
        ws.append([f"Persona Ficticia {i}", f"prueba{i}@example.test", f"8888-{i:04}", datetime(2018, 1, 1+i%28)])
    save(wb, "simple.xlsx")
    wb = Workbook(); wb.remove(wb.active)
    for s, n in [("Atenciones", 3500), ("Seguimiento", 1500)]:
        ws = wb.create_sheet(s)
        ws.append(["Nombres", "Apellidos", "Cédula", "Teléfono", "Correo", "Dirección", "Fecha de nacimiento", "Región", "Oficina Local", "Establecimiento", "Código establecimiento", "Tipo servicio", "Monto", "Fecha registro", "Observaciones"])
        for i in range(n):
            ws.append(["María José" if i%2 else "José Andrés", "Peña Quesada" if i%2 else "Muñoz Gómez", f"1-0999-{i%10000:04}", None if i%17==0 else f"8888-{i%10000:04}", f"prueba{i}@example.test", f"Dirección ficticia {i}", datetime(2018,1,1+i%28), "Región de prueba", "Oficina de prueba", "CEN ficticio", 1100102, "Atención de prueba", float(i*100), datetime(2026,1,1), "La madre María Rodríguez llamó al 8888-1234." if i%13==0 else None])
    save(wb, "realista.xlsx")
    wb=Workbook(); ws=wb.active; ws.title="Casos difíciles"
    ws.append(["id","codigo","valor","numero","observaciones","datos","contacto_1","responsable","fecha"])
    for i in range(40):
        ws.append([123456789,1100102,125000,2026,"La madre María Rodríguez llamó al 8888-1234 y escribió a demo@example.test.", "  María José Peña Quesada  " if i%2 else None, ["+506 8888 1234","8888-1234","88881234"][i%3],"José Andrés Muñoz",datetime(2026,1,1)])
    save(wb,"dificil.xlsx")
    wb=Workbook(); ws=wb.active; ws.title="Datos"
    ws.append(["Nombre completo","Monto","Fecha registro","Cálculo"])
    ws.append(["María José Peña Quesada",125000,datetime(2026,1,1),"=B2*2"])
    ws["B2"].number_format='#,##0.00'; ws["C2"].number_format='dd/mm/yyyy'
    ws.merge_cells("A4:B4"); ws["A4"]="Persona Ficticia"
    dv=DataValidation(type="whole",operator="between",formula1=0,formula2=1000000); dv.add(ws["B2"]); ws.add_data_validation(dv)
    ws=wb.create_sheet("Otra hoja"); ws.append(["Nombre completo","Total"]); ws.append(["María José Peña Quesada","=Datos!B2"])
    ws=wb.create_sheet("Oculta"); ws.append(["Correo"]); ws.append(["demo@example.test"]); ws.sheet_state="hidden"
    save(wb,"complejo.xlsx")
    print("Creados 4 libros de prueba completamente sintéticos.")

if __name__ == "__main__": generate()
